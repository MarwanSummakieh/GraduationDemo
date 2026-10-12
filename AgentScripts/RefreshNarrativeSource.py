"""Refresh versioned narrative CSV inputs from the writers' ZIP and XLSX.

Usage: python AgentScripts/RefreshNarrativeSource.py --zip export.zip --workbook matrix.xlsx
Then use Graduation Demo > Narrative > Refresh Catalogue in Unity.
Requires openpyxl. Does not modify the writers' input files.
"""
import argparse
import csv
import io
import zipfile
from pathlib import Path
import openpyxl

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--zip", required=True, type=Path)
parser.add_argument("--workbook", required=True, type=Path)
args = parser.parse_args()
target = Path(__file__).resolve().parent.parent / "Assets/Narrative/Source"
required = {"clues.csv", "statements.csv", "routes.csv", "voice_lines.csv", "spoken_script.csv", "unlock_rules.csv"}
with zipfile.ZipFile(args.zip) as archive:
    entries = [name for name in archive.namelist() if Path(name).name in required]
    if len(entries) != len(required) or {Path(name).name for name in entries} != required:
        raise ValueError("Narrative ZIP must contain each required CSV exactly once")
    staged = {Path(name).name: archive.read(name) for name in entries}
workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
if "Theory Matrix" not in workbook.sheetnames:
    raise ValueError("Workbook has no Theory Matrix sheet")
rows = list(workbook["Theory Matrix"].iter_rows(values_only=True))
if not rows or tuple(rows[0][:2]) != ("Theory ID", "Clue ID"):
    raise ValueError("Unsupported theory matrix header")
output = io.StringIO(newline="")
csv.writer(output).writerows(rows)
staged["theories.csv"] = output.getvalue().encode("utf-8")
target.mkdir(parents=True, exist_ok=True)
for name, data in staged.items():
    (target / name).write_bytes(data)
(target / "Murder_Most_Inconvenient_Unity_Theory_Voice_v3.xlsx").write_bytes(args.workbook.read_bytes())
print(f"Refreshed {len(staged)} source CSV files. Run Unity Refresh Catalogue for validated import.")
